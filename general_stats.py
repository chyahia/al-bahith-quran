# -*- coding: utf-8 -*-
"""
«الإحصاء العام» (القسم 21 من السجل). قراءة فقط.

ثلاثة أقسام:
  1. ترتيب الجذور بحسب ورودها (مع الآيات والسور والمكي والمدني وعدد الصيغ).
  2. جدول السور وملف كل سورة: الأسماء والأفعال والأدوات، ونسبة الأفعال إلى الأسماء، وأكثر جذورها ورودًا،
     والجذور المميِّزة لها (أعلى ورودًا فيها من نصيبها المتوقع).
  3. المكي والمدني: المؤشرات نفسها، والجذور الأكثر ورودًا في كل منهما، والجذور الأميل إلى أحدهما، والخاصة به.

المصادر بالطريقة نفسها في «ملف الجذر»:
  - مواضع الجذور من root_words (أول جذر للموضع، _root_index في root_profile.py).
  - الصيغ من word_lemmas.
  - تصنيف الكلمة (اسم/فعل/أداة) من MASAQ: كل كلمة في MASAQ تُصنَّف بـ classify_word_kind، وهو تعريف
    «الاسم» و«الفعل» في لوحة التبويب 1. الكلمة التي لا تُصنَّف اسمًا ولا فعلًا تُعدّ «حرفًا أو أداة».
    وعدد الكلمات هنا عدد كلمات MASAQ (يختلف عن تقسيم النص بالمسافات في مواضع قليلة).
  - المكي: ترتيب النزول <= 86، كما في «ملف الجذر».

الحساب ثقيل نسبيًا (كل MASAQ)، فيُحسب مرة واحدة ويُحفظ في الذاكرة حتى إعادة تشغيل البرنامج.
"""
import threading
from collections import Counter, defaultdict
from grammar_engine import get_db_connection, classify_word_kind
from root_profile import _root_index

MECCAN_MAX_REV = 86
MIN_ROOT_FOR_LEANING = 10      # «الأميل إلى المكي/المدني»: جذر لا يقل وروده عن هذا
MIN_SURA_ROOT_FOR_CHAR = 3     # «الجذور المميِّزة للسورة»: لا يقل ورود الجذر في السورة عن هذا
TOP_N = 20

_cache = None
_lock = threading.Lock()


def _pct(a, b, digits=1):
    return round(100.0 * a / b, digits) if b else 0.0


def _ratio(a, b, digits=2):
    return round(a / b, digits) if b else None


def _compute(rev_order):
    q = get_db_connection('quran.db')
    names = {int(r['id']): r['name'] for r in q.execute("SELECT id, name FROM suras")}
    ayas_per_sura = Counter()
    for r in q.execute("SELECT sura_id FROM ayas"):
        ayas_per_sura[int(r['sura_id'])] += 1
    by_pos, freq = _root_index(q.cursor())
    lemma_of = {}
    try:
        for s, a, w, lem in q.execute("SELECT sura_id, aya_num, word_num, lemma FROM word_lemmas"):
            lemma_of[(int(s), int(a), int(w))] = lem
    except Exception:
        pass  # قاعدة بلا word_lemmas: عمود «الصيغ» يبقى فارغًا
    q.close()

    def is_meccan(s):
        return rev_order.get(s, 1) <= MECCAN_MAX_REV

    # ---- تصنيف كلمات MASAQ
    m = get_db_connection('MASAQ.db')
    words = Counter()                       # سورة -> عدد الكلمات
    kinds = defaultdict(Counter)            # سورة -> noun/verb/other
    labels = defaultdict(Counter)           # سورة -> تسمية التصنيف (فعل ماضٍ، اسم مرفوع...)
    cur_key, segs = None, []

    def flush():
        if cur_key is None:
            return
        s = cur_key[0]
        words[s] += 1
        k = classify_word_kind(segs)
        if k is None:
            kinds[s]['other'] += 1
        else:
            kinds[s][k[0]] += 1
            labels[s][k[1]] += 1

    for r in m.execute("""SELECT Sura_No, Verse_No, Word_No, Morph_Tag, Case_Mood FROM MASAQ
                          ORDER BY Sura_No, Verse_No, Word_No, Segment_No"""):
        key = (int(r['Sura_No']), int(r['Verse_No']), int(r['Word_No']))
        if key != cur_key:
            flush()
            cur_key, segs = key, []
        segs.append((r['Morph_Tag'], r['Case_Mood']))
    flush()
    m.close()

    # ---- الجذور
    root_ayas, root_suras = defaultdict(set), defaultdict(set)
    root_meccan, root_lemmas = Counter(), defaultdict(set)
    sura_roots = defaultdict(Counter)
    for (s, a, w), r in by_pos.items():
        root_ayas[r].add((s, a))
        root_suras[r].add(s)
        if is_meccan(s):
            root_meccan[r] += 1
        lem = lemma_of.get((s, a, w))
        if lem:
            root_lemmas[r].add(lem)
        sura_roots[s][r] += 1

    roots = []
    for r, c in freq.items():
        roots.append({'root': r, 'count': c, 'ayahs': len(root_ayas[r]), 'suras': len(root_suras[r]),
                      'meccan': root_meccan[r], 'medinan': c - root_meccan[r],
                      'lemmas': len(root_lemmas[r])})
    roots.sort(key=lambda x: (-x['count'], x['root']))
    for i, x in enumerate(roots, 1):
        x['rank'] = i

    # ---- جدول السور
    suras = []
    for s in sorted(names):
        k = kinds[s]
        suras.append({
            'id': s, 'name': names[s], 'type': 'مكية' if is_meccan(s) else 'مدنية', 'rev': rev_order.get(s),
            'ayas': ayas_per_sura[s], 'words': words[s],
            'nouns': k['noun'], 'verbs': k['verb'], 'other': k['other'],
            'verb_noun': _ratio(k['verb'], k['noun']),
            'verbs_pct': _pct(k['verb'], words[s]),
            'root_positions': sum(sura_roots[s].values()), 'distinct_roots': len(sura_roots[s]),
        })

    # ---- المكي والمدني
    def side(meccan):
        ss = [x for x in suras if (x['type'] == 'مكية') == meccan]
        cnt = Counter()
        for x in ss:
            cnt.update(sura_roots[x['id']])
        tot = lambda f: sum(x[f] for x in ss)
        return ss, cnt, {
            'suras': len(ss), 'ayas': tot('ayas'), 'words': tot('words'),
            'nouns': tot('nouns'), 'verbs': tot('verbs'), 'other': tot('other'),
            'root_positions': sum(cnt.values()), 'distinct_roots': len(cnt),
        }

    _, cm, sm = side(True)
    _, cd, sd = side(False)
    for st in (sm, sd):
        st['verb_noun'] = _ratio(st['verbs'], st['nouns'])
        st['verbs_pct'] = _pct(st['verbs'], st['words'])
        st['words_per_aya'] = _ratio(st['words'], st['ayas'], 1)
    sm['exclusive_roots'] = sum(1 for r in cm if r not in cd)
    sd['exclusive_roots'] = sum(1 for r in cd if r not in cm)

    pm, pd = sm['root_positions'] or 1, sd['root_positions'] or 1

    def top(cnt, p):
        return [{'root': r, 'count': c, 'per_1000': round(1000.0 * c / p, 1)}
                for r, c in sorted(cnt.items(), key=lambda x: (-x[1], x[0]))[:TOP_N]]

    # الميل: نسبة معدل الورود في المكي إلى معدله في المدني (لكل ألف موضع جذر)، مع تمهيد 0.5
    # حتى لا يُقسم على صفر. ويُشترط حدّ أدنى لورود الجذر كي لا تتصدّر الجذور النادرة.
    leaning = []
    for r, c in freq.items():
        if c < MIN_ROOT_FOR_LEANING:
            continue
        a, b = cm.get(r, 0), cd.get(r, 0)
        leaning.append({'root': r, 'meccan': a, 'medinan': b, 'count': c,
                        'lean': round(((a + 0.5) / pm) / ((b + 0.5) / pd), 2)})
    lean_m = sorted([x for x in leaning if x['lean'] > 1], key=lambda x: (-x['lean'], -x['count'], x['root']))[:TOP_N]
    lean_d = sorted([x for x in leaning if x['lean'] < 1], key=lambda x: (x['lean'], -x['count'], x['root']))[:TOP_N]
    for x in lean_d:
        x['lean_inv'] = round(1 / x['lean'], 2) if x['lean'] else None

    def exclusive(cnt, other):
        return [{'root': r, 'count': c} for r, c in sorted(cnt.items(), key=lambda x: (-x[1], x[0]))
                if r not in other][:TOP_N]

    k_all = Counter()
    for s in kinds:
        k_all.update(kinds[s])
    overview = {
        'suras': len(names), 'ayas': sum(ayas_per_sura.values()), 'words': sum(words.values()),
        'nouns': k_all['noun'], 'verbs': k_all['verb'], 'other': k_all['other'],
        'root_positions': len(by_pos), 'distinct_roots': len(freq),
        'distinct_lemmas': len(set(lemma_of.values())),
    }

    return {
        'overview': overview,
        'roots': roots,
        'suras': suras,
        'mm': {
            'meccan': sm, 'medinan': sd,
            'top_meccan': top(cm, pm), 'top_medinan': top(cd, pd),
            'lean_meccan': lean_m, 'lean_medinan': lean_d,
            'excl_meccan': exclusive(cm, cd), 'excl_medinan': exclusive(cd, cm),
            'min_root_for_leaning': MIN_ROOT_FOR_LEANING,
        },
        # للملف التفصيلي لكل سورة (لا يُرسل كاملًا إلى الواجهة)
        '_sura_roots': sura_roots, '_labels': labels, '_freq': freq, '_n_total': len(by_pos),
    }


def get_general_stats(rev_order, refresh=False):
    global _cache
    with _lock:
        if _cache is None or refresh:
            _cache = _compute(rev_order)
        return _cache


def public_stats(rev_order, refresh=False):
    d = get_general_stats(rev_order, refresh)
    return {k: v for k, v in d.items() if not k.startswith('_')}


def sura_profile(sura_id, rev_order):
    d = get_general_stats(rev_order)
    s = int(sura_id)
    row = next((x for x in d['suras'] if x['id'] == s), None)
    if row is None:
        return {'error': 'سورة غير موجودة'}
    cnt = d['_sura_roots'].get(s, Counter())
    pos_s = sum(cnt.values()) or 1
    freq, n_total = d['_freq'], d['_n_total']
    top_roots = [{'root': r, 'count': c, 'pct': _pct(c, pos_s), 'root_freq': freq[r]}
                 for r, c in sorted(cnt.items(), key=lambda x: (-x[1], x[0]))[:TOP_N]]
    # المميِّزة: نسبة نصيب الجذر من مواضع السورة إلى نصيبه من مواضع القرآن كله
    char = []
    for r, c in cnt.items():
        if c >= MIN_SURA_ROOT_FOR_CHAR:
            char.append({'root': r, 'count': c, 'root_freq': freq[r],
                         'share_of_root': _pct(c, freq[r]),
                         'keyness': round((c / pos_s) / (freq[r] / n_total), 1)})
    char.sort(key=lambda x: (-x['keyness'], -x['count'], x['root']))
    lab = d['_labels'].get(s, Counter())
    nouns = [[k, v] for k, v in lab.most_common() if not k.startswith('فعل')]
    verbs = [[k, v] for k, v in lab.most_common() if k.startswith('فعل')]
    return dict(row, top_roots=top_roots, characteristic=char[:TOP_N], noun_labels=nouns, verb_labels=verbs,
                min_sura_root=MIN_SURA_ROOT_FOR_CHAR)


def export_general_stats_excel(rev_order):
    import openpyxl
    from openpyxl.styles import Font, Alignment

    d = get_general_stats(rev_order)
    wb = openpyxl.Workbook()

    def sheet(ws, headers, rows):
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal='center')
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            mx = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(mx + 3, 45)

    sp = lambda r: ' '.join(r)
    ws = wb.active
    ws.title = 'الجذور'
    sheet(ws, ['الترتيب', 'الجذر', 'المواضع', 'الآيات', 'السور', 'مكي', 'مدني', 'الصيغ (اللمات)'],
          [[x['rank'], sp(x['root']), x['count'], x['ayahs'], x['suras'], x['meccan'], x['medinan'], x['lemmas']]
           for x in d['roots']])

    sheet(wb.create_sheet('السور'),
          ['رقم', 'السورة', 'النوع', 'ترتيب النزول', 'الآيات', 'الكلمات', 'الأسماء', 'الأفعال', 'حروف وأدوات',
           'الأفعال/الأسماء', 'نسبة الأفعال %', 'مواضع الجذور', 'جذور مختلفة'],
          [[x['id'], x['name'], x['type'], x['rev'], x['ayas'], x['words'], x['nouns'], x['verbs'], x['other'],
            x['verb_noun'], x['verbs_pct'], x['root_positions'], x['distinct_roots']] for x in d['suras']])

    mm = d['mm']
    ws = wb.create_sheet('المكي والمدني')
    labels = [('suras', 'السور'), ('ayas', 'الآيات'), ('words', 'الكلمات'), ('words_per_aya', 'متوسط كلمات الآية'),
              ('nouns', 'الأسماء'), ('verbs', 'الأفعال'), ('other', 'حروف وأدوات'), ('verb_noun', 'الأفعال/الأسماء'),
              ('verbs_pct', 'نسبة الأفعال %'), ('root_positions', 'مواضع الجذور'), ('distinct_roots', 'جذور مختلفة'),
              ('exclusive_roots', 'جذور خاصة به')]
    sheet(ws, ['المؤشر', 'المكي', 'المدني'], [[t, mm['meccan'][k], mm['medinan'][k]] for k, t in labels])

    sheet(wb.create_sheet('الأميل إلى المكي'), ['الجذر', 'في المكي', 'في المدني', 'المجموع', 'الميل'],
          [[sp(x['root']), x['meccan'], x['medinan'], x['count'], x['lean']] for x in mm['lean_meccan']])
    sheet(wb.create_sheet('الأميل إلى المدني'), ['الجذر', 'في المكي', 'في المدني', 'المجموع', 'الميل'],
          [[sp(x['root']), x['meccan'], x['medinan'], x['count'], x['lean_inv']] for x in mm['lean_medinan']])
    return wb
