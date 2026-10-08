# -*- coding: utf-8 -*-
"""
«ملف الجذر»: كل ما يتعلق بجذر واحد (القسم 20 من السجل). قراءة فقط.
يعتمد على: root_words وroots وayas وword_lemmas في quran.db، وMASAQ.db لتصنيف الاسم/الفعل
(بالمنطق نفسه في لوحة التبويب 1: مواءمة الترقيم، تقسيم الكلمة المزدوجة، classify_word_kind).
"""
import sqlite3
from collections import Counter, defaultdict
from grammar_engine import (
    get_db_connection, adjust_word_no_for_masaq, split_shared_word_segments, classify_word_kind,
)

POS_AR = {
    'V': 'فعل', 'N': 'اسم', 'PN': 'علم', 'ADJ': 'صفة', 'IMPN': 'اسم فعل', 'PRON': 'ضمير',
    'DEM': 'اسم إشارة', 'REL': 'اسم موصول', 'T': 'ظرف زمان', 'LOC': 'ظرف مكان',
    'INTG': 'استفهام', 'COND': 'شرط', 'P': 'حرف جر', 'NEG': 'نفي',
}
MIN_FREQ_FOR_ASSOC = 5     # لا تُحسب «قوة الارتباط» لجذر أقل ورودًا من هذا
MIN_PAIRS_FOR_ASSOC = 3    # ولا لمصاحبة أقل من هذا


def _root_index(qcur):
    """كل مواضع الجذور: (سورة، آية، كلمة) -> الجذر المنظَّف، وعدد ورود كل جذر."""
    by_pos, freq = {}, Counter()
    for s, a, loc, root in qcur.execute("""SELECT rw.sura_id, rw.aya_num, rw.word_location, r.arabic_trilateral
                                          FROM root_words rw JOIN roots r ON rw.root_id = r.id"""):
        try:
            k = (int(s), int(a), int(loc.split(':')[2]))
        except Exception:
            continue
        clean = root.replace(' ', '')
        if k not in by_pos:
            by_pos[k] = clean
            freq[clean] += 1
    return by_pos, freq


def build_root_profile(root, window=3, rev_order=None, sura_names=None):
    clean = (root or '').replace(' ', '')
    rev_order = rev_order or {}
    q = get_db_connection('quran.db')
    qcur = q.cursor()
    by_pos, freq = _root_index(qcur)
    positions = sorted(k for k, r in by_pos.items() if r == clean)
    if not positions:
        q.close()
        return {'root': clean, 'total': 0}

    # نصوص الآيات (للصور المكتوبة)
    ayas = {}
    for s, a, tash, uth in qcur.execute("SELECT sura_id, aya_num, text_tashkeel, text_uthmani FROM ayas"):
        ayas[(int(s), int(a))] = (tash.split(' '), uth.split(' '))

    # اللمات
    lemmas = {}
    for s, a, w, lem, pos in qcur.execute("SELECT sura_id, aya_num, word_num, lemma, pos FROM word_lemmas"):
        lemmas[(int(s), int(a), int(w))] = (lem, pos)
    q.close()

    # تصنيف الاسم/الفعل من MASAQ
    m = get_db_connection('MASAQ.db')
    ayakeys = {(s, a) for s, a, _ in positions}
    segs_by_word = defaultdict(list)
    keys = list(ayakeys)
    for i in range(0, len(keys), 400):
        chunk = keys[i:i + 400]
        ph = ','.join('(?,?)' for _ in chunk)
        for r in m.execute(f"""SELECT Sura_No, Verse_No, Word_No, Segment_No, Word, Morph_Tag, Case_Mood FROM MASAQ
                               WHERE (Sura_No, Verse_No) IN ({ph})
                               ORDER BY Sura_No, Verse_No, Word_No, Segment_No""",
                           [v for t in chunk for v in t]):
            segs_by_word[(r['Sura_No'], r['Verse_No'], r['Word_No'])].append(
                (r['Segment_No'], r['Morph_Tag'], r['Case_Mood'], r['Word']))
    m.close()

    kind_of, word_of = {}, {}
    for s, a, w in positions:
        mw = adjust_word_no_for_masaq(s, a, w)
        segs = segs_by_word.get((s, a, mw), []) if mw is not None else []
        segs = split_shared_word_segments(s, a, w, mw, segs) if segs else segs
        kind_of[(s, a, w)] = classify_word_kind([(t, cm) for _, t, cm, _ in segs]) if segs else None
        # الكلمة نفسها من MASAQ (مواءَمة الترقيم ومتحقَّق منها بأداة «ربط الجذور»)، لا من تقسيم نص الآية
        # بالمسافات: ترقيم المدونة لا يطابق كلمات النص المشكول دائمًا («يَا أَيُّهَا» كلمتان فيه، واحدة في المدونة).
        word_of[(s, a, w)] = next((x[3] for x in segs if x[3]), None)

    # ---- الرأس
    suras = sorted({s for s, _, _ in positions})
    meccan = sum(1 for s, _, _ in positions if rev_order.get(s, 1) <= 86)
    first_mushaf = positions[0]
    first_rev = min(positions, key=lambda p: (rev_order.get(p[0], 999), p[1], p[2]))

    def written(p):
        if word_of.get(p):
            return (word_of[p], '')
        tash, uth = ayas.get((p[0], p[1]), ([], []))
        i = p[2] - 1
        return (tash[i] if 0 <= i < len(tash) else '', uth[i] if 0 <= i < len(uth) else '')

    # ---- اللمات
    lem_groups = defaultdict(list)
    for p in positions:
        lem, pos = lemmas.get(p, (None, None))
        lem_groups[(lem, pos)].append(p)
    total = len(positions)
    lemma_rows = []
    for (lem, pos), ps in sorted(lem_groups.items(), key=lambda x: -len(x[1])):
        forms = Counter(written(p)[0] for p in ps)
        kinds = Counter((kind_of.get(p) or ('', 'غير مصنّف'))[1] for p in ps)
        lemma_rows.append({
            'lemma': lem or 'غير محددة', 'pos': POS_AR.get(pos, pos or ''),
            'count': len(ps), 'percent': round(100 * len(ps) / total, 1),
            'forms': [{'form': f, 'count': c} for f, c in forms.most_common(6) if f],
            'kinds': [{'kind': k, 'count': c} for k, c in kinds.most_common()],
            'positions': [[p[0], p[1], p[2], written(p)[0]] for p in ps],
        })

    # ---- الأسماء والأفعال (كلوحة التبويب 1)
    nouns, verbs = Counter(), Counter()
    for p in positions:
        k = kind_of.get(p)
        if k:
            (nouns if k[0] == 'noun' else verbs)[k[1]] += 1

    # ---- التوزيع على السور
    per_sura = Counter(s for s, _, _ in positions)
    top_suras = [{'sura_id': s, 'name': (sura_names or {}).get(s, str(s)), 'count': c}
                 for s, c in per_sura.most_common(10)]

    # ---- الجذور المصاحبة في نافذة من الكلمات
    n_total = sum(freq.values())
    pairs = Counter()
    for s, a, w in positions:
        for d in range(-window, window + 1):
            if d == 0:
                continue
            r = by_pos.get((s, a, w + d))
            if r and r != clean:
                pairs[r] += 1
    co = []
    for r, o in pairs.items():
        assoc = None
        if freq[r] >= MIN_FREQ_FOR_ASSOC and o >= MIN_PAIRS_FOR_ASSOC:
            expected = total * (2 * window) * freq[r] / n_total
            assoc = round(o / expected, 2) if expected else None
        co.append({'root': r, 'count': o, 'root_freq': freq[r], 'assoc': assoc})
    co_by_count = sorted(co, key=lambda x: (-x['count'], x['root']))[:20]
    co_by_assoc = sorted([x for x in co if x['assoc'] is not None], key=lambda x: -x['assoc'])[:20]

    return {
        'root': clean, 'total': total,
        'ayahs': len({(s, a) for s, a, _ in positions}),
        'suras': len(suras), 'meccan': meccan, 'medinan': total - meccan,
        'first_mushaf': {'sura': first_mushaf[0], 'name': (sura_names or {}).get(first_mushaf[0], ''),
                         'aya': first_mushaf[1], 'word': written(first_mushaf)[0]},
        'first_revelation': {'sura': first_rev[0], 'name': (sura_names or {}).get(first_rev[0], ''),
                             'aya': first_rev[1], 'word': written(first_rev)[0]},
        'lemmas': lemma_rows,
        'nouns': {'total': sum(nouns.values()), 'details': nouns.most_common()},
        'verbs': {'total': sum(verbs.values()), 'details': verbs.most_common()},
        'top_suras': top_suras,
        'positions_suras': [s for s, _, _ in positions],
        'window': window,
        'co_by_count': co_by_count, 'co_by_assoc': co_by_assoc,
    }


def root_profile_ayahs(positions):
    """نصوص الآيات لقائمة مواضع [(سورة، آية، كلمة، الكلمة)]، والكلمة تُعرض بعد الآية (بلا تلوين)."""
    q = get_db_connection('quran.db')
    out = []
    names = {r['id']: r['name'] for r in q.execute("SELECT id, name FROM suras")}
    for p in positions:
        s, a, w = int(p[0]), int(p[1]), int(p[2])
        word = p[3] if len(p) > 3 else ''
        r = q.execute("SELECT text_tashkeel FROM ayas WHERE sura_id = ? AND aya_num = ?", (s, a)).fetchone()
        if r:
            out.append({'sura': s, 'sura_name': names.get(s, str(s)), 'aya': a, 'word_no': w,
                        'word': word, 'text': r['text_tashkeel']})
    q.close()
    return out


def root_pair_stats(root1, root2, window=3):
    """اجتماع جذرين: عدد الآيات التي فيها الاثنان، وعدد مواضع الأول التي يقع الثاني في نافذتها."""
    r1, r2 = (root1 or '').replace(' ', ''), (root2 or '').replace(' ', '')
    q = get_db_connection('quran.db')
    by_pos, freq = _root_index(q.cursor())
    q.close()
    ayas1 = {(s, a) for (s, a, w), r in by_pos.items() if r == r1}
    ayas2 = {(s, a) for (s, a, w), r in by_pos.items() if r == r2}
    near = 0
    for (s, a, w), r in by_pos.items():
        if r != r1:
            continue
        if any(by_pos.get((s, a, w + d)) == r2 for d in range(-window, window + 1) if d):
            near += 1
    return {'shared_ayahs': len(ayas1 & ayas2), 'near': near, 'window': window,
            'freq1': freq.get(r1, 0), 'freq2': freq.get(r2, 0)}


# ---------------------------------------------------------------------------
# تصدير ملف الجذر (أو مقارنة جذرين) إلى Word (RTL، بطريقة export.py)
# ---------------------------------------------------------------------------
def create_root_profile_docx(d, chart_b64=None, d2=None, pair=None):
    import io, base64
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    GREEN = RGBColor(27, 94, 32)
    RLM = '\u200F'
    doc = Document()
    sect = doc.sections[0]._sectPr
    if sect.find(qn('w:bidi')) is None:
        sect.append(OxmlElement('w:bidi'))
    theme_font_lang = doc.settings.element.find(qn('w:themeFontLang'))
    if theme_font_lang is not None:
        theme_font_lang.set(qn('w:bidi'), 'ar-SA')

    def rtl_par(p):
        # في الفقرة ثنائية الاتجاه تكون المحاذاة «منطقية»: LEFT = بداية السطر = اليمين في العربية
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement('w:bidi'))
        mark_rPr = OxmlElement('w:rPr')
        mark_rPr.append(OxmlElement('w:rtl'))
        pPr.append(mark_rPr)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def run(p, text, size=12, bold=False, color=None, font=None):
        r = p.add_run(RLM + str(text) + RLM)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.rtl = True
        if color:
            r.font.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        f = font or 'Arial'
        for k in ('w:cs', 'w:ascii', 'w:hAnsi'):
            rFonts.set(qn(k), f)
        rFonts.set(qn('w:hint'), 'cs')
        szCs = OxmlElement('w:szCs')
        szCs.set(qn('w:val'), str(int(size * 2)))
        rPr.append(szCs)

    def para(text='', size=12, bold=False, color=None, font=None):
        p = doc.add_paragraph()
        rtl_par(p)
        if text:
            run(p, text, size, bold, color, font)
        return p

    def table(headers, rows, fonts=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i].paragraphs[0]
            rtl_par(c)
            run(c, h, 11, True)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                c = cells[i].paragraphs[0]
                rtl_par(c)
                f = (fonts or {}).get(i)
                run(c, v, 14 if f else 11, False, None, f)
        doc.add_paragraph()

    def chart(b64):
        if b64 and ',' in b64:
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(io.BytesIO(base64.b64decode(b64.split(',', 1)[1])), width=Inches(6.0))
            except Exception:
                pass

    spaced = lambda r: ' '.join(r)

    def profile_body(x, with_chart):
        para(f"ملف الجذر «{spaced(x['root'])}»", 16, True, GREEN)
        para(f"عدد المواضع: {x['total']}  |  عدد الآيات: {x['ayahs']}  |  عدد السور: {x['suras']}  |  "
             f"مكي/مدني: {x['meccan']}/{x['medinan']}  |  عدد الصيغ: {len(x['lemmas'])}", 12)
        fm, fr = x['first_mushaf'], x['first_revelation']
        para(f"أول ورود في المصحف: {fm['name']} {fm['aya']} ({fm['word']})  —  "
             f"وفي النزول: {fr['name']} {fr['aya']} ({fr['word']})", 12)
        para('صيغ الجذر (اللمات)', 14, True, GREEN)
        table(['اللمة', 'النوع', 'العدد', 'النسبة', 'أكثر الصور ورودًا'],
              [[l['lemma'], l['pos'], l['count'], f"{l['percent']}%",
                '، '.join(f"{f['form']} ({f['count']})" for f in l['forms'])] for l in x['lemmas']],
              fonts={0: 'Traditional Arabic'})
        para(f"الأسماء ({x['nouns']['total']}) والأفعال ({x['verbs']['total']})", 14, True, GREEN)
        table(['التصنيف', 'العدد'], [[k, v] for k, v in x['nouns']['details'] + x['verbs']['details']])
        para('التوزيع', 14, True, GREEN)
        if with_chart:
            chart(with_chart)
        table(['السورة', 'عدد المواضع'], [[s_['name'], s_['count']] for s_ in x['top_suras']])
        para(f"الجذور المصاحبة (نافذة {x['window']} كلمات)", 14, True, GREEN)
        para('قوة الارتباط: نسبة الاجتماع الفعلي إلى المتوقع من شيوع الجذرين (أكبر من 1 = أكثر من المتوقع).', 10)
        table(['الجذر', 'عدد المصاحبة', 'ورود الجذر', 'قوة الارتباط'],
              [[spaced(c['root']), c['count'], c['root_freq'], '—' if c['assoc'] is None else c['assoc']]
               for c in x['co_by_count']])
        if x.get('co_by_assoc'):
            para('الأقوى ارتباطًا', 12, True, GREEN)
            table(['الجذر', 'قوة الارتباط', 'عدد المصاحبة', 'ورود الجذر'],
                  [[spaced(c['root']), c['assoc'], c['count'], c['root_freq']] for c in x['co_by_assoc']])

    if d2:
        para(f"الباحث في القرآن الكريم — مقارنة الجذرين «{spaced(d['root'])}» و«{spaced(d2['root'])}»", 18, True, GREEN)
        table(['المؤشر', spaced(d['root']), spaced(d2['root'])], [
            ['عدد المواضع', d['total'], d2['total']], ['عدد الآيات', d['ayahs'], d2['ayahs']],
            ['عدد السور', d['suras'], d2['suras']], ['مكي', d['meccan'], d2['meccan']],
            ['مدني', d['medinan'], d2['medinan']], ['عدد الصيغ', len(d['lemmas']), len(d2['lemmas'])],
            ['الأسماء', d['nouns']['total'], d2['nouns']['total']], ['الأفعال', d['verbs']['total'], d2['verbs']['total']],
        ])
        if pair:
            para(f"اجتماع الجذرين: في {pair['shared_ayahs']} آية، ويقع «{spaced(d2['root'])}» في نافذة "
                 f"{pair['window']} كلمات حول {pair['near']} موضعًا من مواضع «{spaced(d['root'])}».", 12)
        chart(chart_b64)
        doc.add_page_break()
        profile_body(d, None)
        doc.add_page_break()
        profile_body(d2, None)
    else:
        para('الباحث في القرآن الكريم', 18, True, GREEN)
        profile_body(d, chart_b64)

    para('المصادر: مواضع الجذور من QUL، واللمات من المدونة القرآنية (بالرسم العثماني)، '
         'وتصنيف الاسم/الفعل من MASAQ.', 9)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out
